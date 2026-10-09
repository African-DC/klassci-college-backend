"""Vérifie les champs et le format éditable de la liste Word de classe."""

from datetime import date
from io import BytesIO

from docx import Document

from app.services.class_roster_service import _word_roster_document


def test_word_roster_has_requested_columns_and_values():
    payload = {
        "class_name": "6ème A",
        "academic_year_name": "2026-2027",
        "students": [{
            "enrollment_number": "RST-001",
            "last_name": "KOUAME",
            "first_name": "Amenan",
            "genre": "F",
            "birth_date": date(2013, 4, 2),
            "birth_place": "Bouaké",
            "nationality": "Ivoirienne",
            "is_repeater": False,
            "assignment_status": "affecte",
        }],
    }
    document = Document(BytesIO(_word_roster_document(payload)))
    assert document.sections[0].page_width > document.sections[0].page_height
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == [
        "N°", "Matricule", "Nom", "Prénoms", "Sexe", "Date de naissance",
        "Lieu de naissance", "Nationalité", "Qualité", "Statut",
    ]
    assert [c.text for c in table.rows[1].cells] == [
        "1", "RST-001", "KOUAME", "Amenan", "F", "02/04/2013",
        "Bouaké", "Ivoirienne", "Non redoublant", "Affecté",
    ]


def test_word_roster_leaves_missing_fields_blank():
    data = {
        "class_name": "5ème B",
        "academic_year_name": "2026-2027",
        "students": [{
            "enrollment_number": None,
            "last_name": "YAO",
            "first_name": "N.",
            "genre": None,
            "birth_date": None,
            "birth_place": None,
            "nationality": None,
            "is_repeater": None,
            "assignment_status": None,
        }],
    }
    row = Document(BytesIO(_word_roster_document(data))).tables[0].rows[1]
    assert [c.text for c in row.cells][4:] == ["", "", "", "", "", ""]
