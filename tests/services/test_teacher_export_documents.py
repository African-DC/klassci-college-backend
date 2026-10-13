"""Non-régression : export collège (pas de logique universitaire BTS/LMD)."""
import io

import pytest
from fastapi import HTTPException

from app.services.teacher_export_documents import (
    compose,
    generate_excel,
    generate_word,
    page_landscape,
    validate_options,
)


def teacher(**values):
    return {"name": "KOUAME Aya", "speciality": "Mathématiques", "subject": "Mathématiques",
            "class": "6e A", "level": "6e", "series": "—", "contract": "Permanent",
            "phone": "0700000000", "email": "aya@example.org", **values}


def test_liste_simple_portrait_une_colonne_et_rubrique():
    cols = validate_options(
        view="assignments", columns=["name"], grouping="subject", orientation="auto",
        class_ids=[], level_ids=[], subject_ids=[], series_ids=[],
    )
    assert cols == ["name"]
    assert page_landscape(cols, "auto") is False
    groups = compose([teacher()], cols, "subject", "assignments")
    assert groups[0][0] == "Mathématiques"


def test_colonne_matiere_masquee_et_regroupement_absent_refuse_le_filtre_matiere():
    with pytest.raises(HTTPException) as exc:
        validate_options(view="assignments", columns=["name"], grouping="none",
                         orientation="portrait", class_ids=[], level_ids=[],
                         subject_ids=[12], series_ids=[])
    assert exc.value.status_code == 422


def test_repertoire_ne_devine_pas_des_affectations_depuis_specialite():
    with pytest.raises(HTTPException):
        validate_options(view="directory", columns=["name"], grouping="subject",
                         orientation="auto", class_ids=[], level_ids=[],
                         subject_ids=[], series_ids=[])


def test_portrait_choisi_explicitement_refuse_le_debordement():
    with pytest.raises(HTTPException) as exc:
        page_landscape(["name", "email", "speciality", "subject", "class"], "portrait")
    assert exc.value.status_code == 422


def test_liste_nom_seul_evite_les_repetitions_mais_le_detail_conserve_les_affectations():
    rows = [teacher(), dict(teacher(), **{"class": "6e B"})]
    groups = compose(rows, ["name"], "subject", "assignments")
    assert len(groups[0][1]) == 1
    detail = compose(rows, ["name", "class"], "subject", "assignments")
    assert len(detail[0][1]) == 2
    brut = compose(rows, ["name"], "none", "assignments")
    assert len(brut[0][1]) == 2


def test_deux_matieres_homonymes_ne_sont_pas_fusionnees():
    first = dict(teacher(), subject_id=12)
    second = dict(teacher(), subject_id=22)
    groups = compose([first, second], ["name"], "subject", "assignments")
    assert len(groups) == 2
    assert {g[0] for g in groups} == {"Mathématiques · Réf. #12", "Mathématiques · Réf. #22"}


def test_excel_word_sont_de_vrais_office_files_et_identite_est_conservee():
    rows = [teacher()]
    groups = compose(rows, ["name", "subject"], "subject", "assignments")
    school = {"school_name": "COLLEGE TEST", "primary_color": "#0453CB",
              "accent_color": "#F58220", "motto": "Notre ambition"}
    excel = generate_excel(groups, ["name", "subject"], school, "2026-2027")
    word = generate_word(groups, ["name", "subject"], school, "2026-2027", False)
    assert excel[:2] == b"PK"
    assert word[:2] == b"PK"
    from docx import Document
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(excel)).active
    assert ws["A1"].value == "COLLEGE TEST"
    doc = Document(io.BytesIO(word))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "COLLEGE TEST" in text
    assert "2026-2027" in text
    assert doc.tables[0].rows[1].cells[0].text == "KOUAME Aya"
