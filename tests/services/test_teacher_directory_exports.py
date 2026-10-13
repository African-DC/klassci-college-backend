"""Tests de l'export enseignants : colonnes, filtres cachés, PDF/Word/Excel."""

from __future__ import annotations

from io import BytesIO

import pytest
from pydantic import ValidationError

from app.schemas.teacher_directory import TeacherExportOptions, page_orientation
from app.services.teacher_directory_documents import render_excel, render_word
from app.services.teacher_directory_service import compose_groups


def _row(*, teacher="1", subject="1", subject_name="Mathématiques", klass="6e A"):
    return {
        "name": "KONÉ Aminata", "speciality": "Mathématiques",
        "subject": subject_name, "class": klass, "level": "6e", "series": "—",
        "contract": "Vacataire", "phone": "01020304", "email": "contact@example.org",
        "_teacher_id": teacher, "_subject_id": subject,
        "_class_id": "101", "_level_id": "4", "_series_id": "",
    }


def _directory(columns=None):
    columns = columns or ["name"]
    return {
        "year": "2026-2027",
        "groups": [{"title": "Mathématiques", "rows": [["KONÉ Aminata"]]}],
        "columns": columns,
        "headers": ["Nom et prénoms"],
        "count": 1,
        "teachers_count": 1,
        "group_by": "Par matière",
    }


def test_nom_seul_sans_rubriques_et_filtres_matiere_est_licite():
    options = TeacherExportOptions(
        columns=["name"], group_by="none", subject_ids=[4, 6],
        class_ids=[12], orientation="portrait"
    )
    result = compose_groups([_row(), _row(subject="2")], options)
    assert len(result) == 1
    assert result[0]["rows"] == [["KONÉ Aminata"]]
    assert page_orientation(options) == "portrait"


def test_matiere_homonyme_ne_fusionne_pas_deux_references():
    options = TeacherExportOptions(columns=["name"], group_by="subject")
    groups = compose_groups(
        [_row(subject="42"), _row(subject="99"), _row(subject="42")], options
    )
    assert len(groups) == 2
    assert {group["title"] for group in groups} == {
        "Mathématiques · Réf. #42", "Mathématiques · Réf. #99"
    }
    assert all(len(group["rows"]) == 1 for group in groups)


def test_rubrique_remplace_colonne_repetitive():
    options = TeacherExportOptions(columns=["name"], group_by="class")
    group = compose_groups([_row(), _row(subject="2")], options)
    assert group[0]["title"] == "6e A"
    assert group[0]["rows"] == [["KONÉ Aminata"]]


def test_colonnes_invalides_et_choix_trop_larges_sont_refuses():
    with pytest.raises(ValidationError):
        TeacherExportOptions(columns=["name", "unknown"])
    with pytest.raises(ValidationError):
        TeacherExportOptions(columns=["name", "name"])
    with pytest.raises(ValueError):
        page_orientation(TeacherExportOptions(columns=["name", "subject", "email"], orientation="portrait"))


def test_excel_identite_tenant_et_formule_ne_sexecutent_pas():
    from openpyxl import load_workbook

    directory = _directory()
    directory["groups"][0]["rows"] = [["=HYPERLINK(\"https://invalid.test\", \"click\")"]]
    school = {
        "school_name": "Collège de l'Exemple",
        "primary_color": "#224466", "accent_color": "#DD7700",
        "address": "Bouaké", "logo_url": None,
    }
    data = render_excel(directory, school, TeacherExportOptions(columns=["name"]))
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "Collège de l'Exemple"
    assert ws["A3"].value == "LISTE DES ENSEIGNANTS"
    assert any(cell.value and str(cell.value).startswith("'=HYPERLINK")
               for row in ws for cell in row)
    assert ws.page_setup.orientation == "portrait"


def test_word_reutilise_nom_ecole_orientation_et_colonnes():
    from docx import Document
    from docx.enum.section import WD_ORIENT

    school = {"school_name": "Collège de l'Exemple", "logo_url": None,
              "primary_color": "#224466", "accent_color": "#DD7700"}
    data = render_word(_directory(), school, TeacherExportOptions(columns=["name"]))
    doc = Document(BytesIO(data))
    assert doc.sections[0].orientation == WD_ORIENT.PORTRAIT
    paragraphs = [p.text for p in doc.paragraphs]
    assert "LISTE DES ENSEIGNANTS" in paragraphs
    assert "Collège de l'Exemple" in doc.tables[0].cell(0, 1).text
    assert "KONÉ Aminata" in doc.tables[1].cell(1, 0).text


def test_pdf_reutilise_entete_premium_ecole_sans_couleur_hardcode(monkeypatch):
    weasyprint = pytest.importorskip("weasyprint")
    seen = {}

    class FakeHTML:
        def __init__(self, *, string):
            seen["html"] = string

        def write_pdf(self):
            return b"%PDF-FAKE"

    monkeypatch.setattr(weasyprint, "HTML", FakeHTML)
    from app.services.teacher_directory_documents import render_pdf

    school = {"school_name": "Collège de l'Exemple", "logo_url": None,
              "primary_color": "#224466", "accent_color": "#DD7700"}
    pdf = render_pdf(_directory(), school, TeacherExportOptions(columns=["name"]))
    assert pdf == b"%PDF-FAKE"
    assert "Collège de l&#x27;Exemple" in seen["html"] or "Collège de l'Exemple" in seen["html"]
    assert "#224466" in seen["html"]
    assert "LISTE DES ENSEIGNANTS" in seen["html"]
    assert "KONÉ Aminata" in seen["html"]
