"""Couverture de la composition, sans nécessiter une base de données."""

from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.services.teacher_exports import (
    Selection,
    choose_orientation,
    compose,
    render_docx,
    render_xlsx,
)


def choice(**kwargs):
    return Selection(
        year_id=kwargs.pop("year_id", None),
        class_ids=kwargs.pop("class_ids", []),
        level_ids=[],
        series_ids=[],
        subject_ids=[],
        columns=kwargs.pop("columns", ["name"]),
        group=kwargs.pop("group", "speciality"),
        orientation=kwargs.pop("orientation", "auto"),
        mode=kwargs.pop("mode", "directory"),
        **kwargs,
    )


def sample():
    return [
        {
            "name": "KONE",
            "speciality": "Mathématiques",
            "_id": "1",
            "subject": "Maths",
            "class": "6A",
            "level": "6e",
            "series": "Sans série",
            "_subject_id": "7",
            "_class_id": "3",
            "contract": "Permanent",
            "genre": "M",
            "phone": "",
            "email": "",
        },
        {
            "name": "ADOU",
            "speciality": "Physique",
            "_id": "2",
            "subject": "Maths",
            "class": "5A",
            "level": "5e",
            "series": "Sans série",
            "_subject_id": "13",
            "_class_id": "4",
            "contract": "Vacataire",
            "genre": "F",
            "phone": "",
            "email": "",
        },
    ]


def test_portrait_sur_un_simple_annuaire():
    s = choice()
    s.validate()
    assert choose_orientation(s) == "portrait"
    assert len(compose(sample(), s)) == 2


def test_matiere_homonyme_a_id_distinct():
    s = choice(mode="assignments", year_id=4, group="subject")
    assert len(compose(sample(), s)) == 2
    assert all("matière #" in title for title, _ in compose(sample(), s))


def test_liste_simple_par_matiere_sans_colonne_repetee():
    s = choice(mode="assignments", year_id=4, group="subject")
    assert all(len(value) == 1 for _, rows in compose(sample(), s) for value in rows)


def test_pas_de_perimetre_scolaire_sur_annuaire_global():
    s = choice(class_ids=[1])
    with pytest.raises(HTTPException) as exc:
        s.validate()
    assert exc.value.status_code == 422


def test_pas_de_colonne_inconnue_ou_doublon():
    for columns in [["name", "name"], ["_secret"], []]:
        with pytest.raises(HTTPException):
            choice(columns=columns).validate()


def test_portrait_trop_large_refuse():
    s = choice(columns=["name", "speciality", "phone", "email"], orientation="portrait")
    with pytest.raises(HTTPException):
        choose_orientation(s)


def test_word_a_en_tete_et_noms_editables():
    data = render_docx(
        choice(),
        {
            "school_name": "Collège Démonstration",
            "primary_color": "#0453CB",
            "accent_color": "#F58220",
        },
        sample(),
    )
    from docx import Document

    doc = Document(BytesIO(data))
    assert any("Collège Démonstration" in p.text for p in doc.paragraphs)
    assert len(doc.tables) == 2
    assert "KONE" in "\n".join(
        c.text for table in doc.tables for row in table.rows for c in row.cells
    )


def test_excel_avec_entete_et_anti_injection():
    rows = sample()
    rows[0]["name"] = "=1+2"
    data = render_xlsx(choice(), {"school_name": "Collège Démonstration"}, rows)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert "Collège Démonstration" in [
        cell.value for row in ws for cell in row if isinstance(cell.value, str)
    ]
    assert any(cell.value == "'=1+2" for row in ws for cell in row)
