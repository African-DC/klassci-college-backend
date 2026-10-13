"""Exports enseignants Collège : même sélection pour PDF, Excel et Word.

Le niveau/série provient des classes de l'emploi du temps de l'année choisie.
Une spécialité déclarative n'est JAMAIS une affectation pédagogique.
Lecture seule sur la base du tenant, permissions de l'annuaire enseignant.
"""
from __future__ import annotations

from collections import defaultdict
from html import escape
from io import BytesIO

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.academic import AcademicYear, Class, Level, Series, Subject
from app.models.timetable import TimetableSlot
from app.models.user import TeacherProfile, User

COLUMNS = {
    "name": "Nom et prénoms", "speciality": "Spécialité déclarée",
    "subject": "Matière enseignée", "class": "Classe", "level": "Niveau",
    "series": "Série", "contract": "Type de contrat", "phone": "Téléphone",
    "email": "Email", "genre": "Genre", "source": "Source de l'affectation",
}
GROUPINGS = {"none", "subject", "class", "level", "series", "speciality", "contract"}
VIEWS = {"directory", "assignments"}
CONTRACTS = {"permanent": "Permanent", "vacataire": "Vacataire", "fonctionnaire": "Fonctionnaire"}
MAX_ROWS = 8000


def validate_options(*, view: str, columns: list[str] | None, grouping: str,
                     orientation: str, class_ids: list[int], level_ids: list[int],
                     subject_ids: list[int], series_ids: list[int]) -> list[str]:
    if view not in VIEWS or grouping not in GROUPINGS or orientation not in ("auto", "portrait", "landscape"):
        raise HTTPException(422, "Vue, regroupement ou orientation invalide.")
    cols = columns if columns is not None else (
        ["name", "speciality", "contract", "phone"] if view == "directory"
        else ["name", "subject", "class", "level"]
    )
    if not cols or len(cols) > len(COLUMNS) or len(set(cols)) != len(cols) or any(x not in COLUMNS for x in cols):
        raise HTTPException(422, "Sélectionnez des colonnes valides, sans doublon.")
    if view == "directory" and (grouping in ("subject", "class", "level", "series") or
        class_ids or level_ids or subject_ids or series_ids):
        raise HTTPException(422, "Les filtres pédagogiques exigent la vue Affectations réelles.")
    if view == "assignments" and grouping == "none" and "subject" not in cols and subject_ids:
        raise HTTPException(422, "Affichez la matière ou choisissez un regroupement pour identifier les affectations.")
    for ids in (class_ids, level_ids, subject_ids, series_ids):
        if len(ids) > 100 or len(ids) != len(set(ids)) or any(i < 1 for i in ids):
            raise HTTPException(422, "Sélection multiple invalide (100 éléments maximum).")
    return cols


async def fetch_options(db: AsyncSession) -> dict:
    async def rows(model, cols):
        result = await db.execute(select(*cols).order_by(cols[-1]))
        return [{"id": int(v[0]), "name": str(v[-1])} for v in result.all()]
    return {
        "classes": await rows(Class, [Class.id, Class.name]),
        "levels": await rows(Level, [Level.id, Level.name]),
        "series": await rows(Series, [Series.id, Series.name]),
        "subjects": await rows(Subject, [Subject.id, Subject.name]),
        "years": await rows(AcademicYear, [AcademicYear.id, AcademicYear.name]),
    }


async def load_export(db: AsyncSession, *, view: str, year_id: int | None,
                      class_ids: list[int], level_ids: list[int], series_ids: list[int],
                      subject_ids: list[int], search: str | None = None) -> tuple[list[dict], str]:
    selected_year = None
    if year_id:
        selected_year = (await db.execute(select(AcademicYear).where(AcademicYear.id == year_id))).scalar_one_or_none()
        if selected_year is None:
            raise HTTPException(422, "Année scolaire inconnue.")
    else:
        selected_year = (await db.execute(select(AcademicYear).where(AcademicYear.is_current.is_(True)))).scalar_one_or_none()
    if view == "assignments" and selected_year is None:
        raise HTTPException(422, "Sélectionnez une année scolaire pour les affectations.")

    stmt = select(TeacherProfile, User.email).join(User, User.id == TeacherProfile.user_id).where(
        TeacherProfile.archived_at.is_(None)
    ).order_by(TeacherProfile.last_name, TeacherProfile.first_name, TeacherProfile.id)
    teachers = (await db.execute(stmt)).all()
    base = {}
    for teacher, email in teachers:
        name = " ".join(x for x in (teacher.last_name, teacher.first_name) if x).strip()
        if search and search.casefold() not in (name + " " + (teacher.speciality or "")).casefold():
            continue
        base[teacher.id] = {
            "name": name, "speciality": teacher.speciality or "Non renseignée",
            "contract": CONTRACTS.get(str(getattr(teacher.contract_type, "value", teacher.contract_type) or ""), "Non renseigné"),
            "phone": teacher.phone or "—", "email": email or "—",
            "genre": teacher.genre or "—",
        }
    if view == "directory":
        return list(base.values()), (selected_year.name if selected_year else "")

    stmt = (select(
        TimetableSlot.teacher_id, TimetableSlot.subject_id, Subject.name,
        Class.id, Class.name, Level.id, Level.name, Series.id, Series.name
    ).join(Subject, Subject.id == TimetableSlot.subject_id)
     .join(Class, Class.id == TimetableSlot.class_id)
     .join(Level, Level.id == Class.level_id)
     .outerjoin(Series, Series.id == Class.series_id)
     .where(TimetableSlot.academic_year_id == selected_year.id))
    if class_ids:
        stmt = stmt.where(Class.id.in_(class_ids))
    if level_ids:
        stmt = stmt.where(Level.id.in_(level_ids))
    if series_ids:
        stmt = stmt.where(Series.id.in_(series_ids))
    if subject_ids:
        stmt = stmt.where(Subject.id.in_(subject_ids))
    raw = (await db.execute(stmt)).all()
    result, seen = [], set()
    for teacher_id, subject_id, subject, cls_id, cls, level_id, level, series_id, series in raw:
        if teacher_id not in base:
            continue
        key = (teacher_id, subject_id, cls_id)
        if key in seen:
            continue
        seen.add(key)
        result.append({
            **base[teacher_id], "subject": subject, "subject_id": subject_id,
            "class": cls, "class_id": cls_id, "level": level, "level_id": level_id,
            "series": series or "—", "series_id": series_id,
            "source": "Emploi du temps · année sélectionnée",
        })
        if len(result) > MAX_ROWS:
            raise HTTPException(422, "Plus de 8 000 affectations : affinez les filtres.")
    result.sort(key=lambda x: (x["subject"].casefold(), x["class"].casefold(), x["name"].casefold()))
    return result, selected_year.name


def compose(rows: list[dict], columns: list[str], grouping: str, view: str) -> list[tuple[str, list[dict]]]:
    if grouping == "none":
        # Ne jamais éliminer une affectation différente si l'utilisateur masque
        # la colonne classe. Il a choisi l'affichage, pas une consolidation silencieuse.
        return [("Enseignants" if view == "directory" else "Affectations", rows)]
    grouped = defaultdict(list)
    for row in rows:
        grouped[str(row.get(grouping) or "Non renseigné")].append(row)
    return sorted(grouped.items(), key=lambda item: item[0].casefold())


def page_landscape(columns: list[str], orientation: str) -> bool:
    estimated = sum({
        "name": 50, "subject": 43, "class": 32, "level": 27, "series": 20,
        "speciality": 45, "contract": 28, "phone": 31, "email": 48, "genre": 18, "source": 45,
    }[c] for c in columns)
    if orientation == "portrait" and estimated > 184:
        raise HTTPException(422, "Trop de colonnes pour A4 portrait : retirez des colonnes ou choisissez paysage.")
    if estimated > 268:
        raise HTTPException(422, "Trop de colonnes pour une page A4 : réduisez la sélection.")
    return orientation == "landscape" or (orientation == "auto" and estimated > 184)


def generate_pdf(rows: list[dict], columns: list[str], groups: list[tuple[str, list[dict]]],
                 school: dict, year: str, landscape: bool) -> bytes:
    from weasyprint import HTML

    from app.services.pdf import components as ui
    from app.services.pdf.theme import PDFTheme
    theme = PDFTheme.from_school(school)
    blocks = []
    for title, lines in groups:
        blocks.append(f"<h2 class='group-title'>{escape(title)}</h2>")
        data = [[str(row.get(k, "—")) for k in columns] for row in lines]
        blocks.append(ui.premium_table(
            headers=[COLUMNS[k] for k in columns],
            rows=data, theme=theme, empty_message="Aucun enseignant pour cette sélection.",
        ))
    body = f"""<!doctype html><html lang='fr'><head><meta charset='utf-8'>
        {ui.base_styles(theme, page_size="A4 landscape" if landscape else "A4", margin="12mm")}
        <style>.group-title{{font-size:12px;color:var(--primary);margin:14px 0 5px;break-after:avoid}}
        .group-title + table{{break-before:avoid}}</style></head><body>
        {ui.page_decoration(theme=theme, watermark_text=school.get("school_name") or "")}
        <div class='pdf-page-body'>
        {ui.premium_header(school, theme=theme, doc_type="LISTE DES ENSEIGNANTS", doc_subtitle=f"Année scolaire : {escape(year)}")}
        {ui.meta_banner(f"<strong>{len(rows)} affectation(s) / fiche(s)</strong>", "<span>Document interne — KLASSCI Collège</span>", theme=theme)}
        {''.join(blocks)}
        {ui.premium_footer(school, theme=theme, note="Export du répertoire enseignant — données issues de l'établissement.")}
        </div></body></html>"""
    return HTML(string=body).write_pdf()


def generate_excel(groups: list[tuple[str, list[dict]]], columns: list[str],
                   school: dict, year: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.services.exports._workbook_branding import argb, write_header
    wb = Workbook()
    ws = wb.active
    ws.title = "Enseignants"
    idx = write_header(ws, school, title="LISTE DES ENSEIGNANTS",
                       subtitle=f"Année scolaire : {year}", meta_lines=[], width=len(columns))
    primary = argb(school.get("primary_color"), "#0F3F8C")
    row_index = max(idx, ws.max_row + 1)
    for title, lines in groups:
        ws.cell(row_index, 1, value=title).font = Font(bold=True, color=primary)
        row_index += 1
        for i, key in enumerate(columns, 1):
            cell = ws.cell(row_index, i, value=COLUMNS[key])
            cell.fill = PatternFill("solid", fgColor=primary)
            cell.font = Font(bold=True, color="FFFFFFFF")
        row_index += 1
        for line in lines:
            for i, key in enumerate(columns, 1):
                value = str(line.get(key, "—"))
                if value.startswith(("=", "+", "-", "@", "\t", "\r")):
                    value = "'" + value
                ws.cell(row_index, i, value=value)
            row_index += 1
    for i, col in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(58, max(17, len(COLUMNS[col]) + 6))
    ws.sheet_view.showGridLines = False
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_word(groups: list[tuple[str, list[dict]]], columns: list[str],
                  school: dict, year: str, landscape: bool) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, RGBColor

    from app.services.pdf._helpers import image_bytes
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.6)
    logo = school.get("logo_url")
    if logo:
        try:
            image = image_bytes(logo)
            if image:
                doc.add_picture(BytesIO(image[0]), width=Inches(.8))
        except (OSError, ValueError, TypeError):
            pass
    school_title = doc.add_paragraph()
    school_title.add_run(school.get("school_name") or "Établissement").bold = True
    for value in ("ministry_code", "address", "phone", "email", "motto"):
        if school.get(value):
            doc.add_paragraph(str(school[value]), style="Normal")
    primary = (school.get("primary_color") or "#0F3F8C").lstrip("#")
    try:
        rgb = RGBColor.from_string(primary)
        school_title.runs[0].font.color.rgb = rgb
    except ValueError:
        rgb = RGBColor(15, 63, 140)
    doc.add_heading("LISTE DES ENSEIGNANTS", 0)
    doc.add_paragraph(f"Année scolaire : {year}")
    for title, lines in groups:
        doc.add_heading(title, level=2)
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell in table.rows[0].cells:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), primary.upper() if len(primary) == 6 else "0F3F8C")
            cell._tc.get_or_add_tcPr().append(shading)
        for i, col in enumerate(columns):
            table.rows[0].cells[i].text = COLUMNS[col]
        for line in lines:
            cells = table.add_row().cells
            for i, col in enumerate(columns):
                cells[i].text = str(line.get(col, "—"))
    doc.add_paragraph("Document interne • Généré depuis KLASSCI Collège")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
