"""Service : composition + génération de la liste de classe (PDF)."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import NotFoundError
from app.models.academic import AcademicYear, Class
from app.models.enrollment import Enrollment, EnrollmentStatus
from app.models.user import ParentStudent
from app.services._school_settings_helper import (
    load_school_settings_for_pdf as _get_school_settings_dict,
)
from app.services.pdf import generate_class_roster_pdf


async def _load_class_with_room(db: AsyncSession, class_id: int) -> Class:
    stmt = (
        select(Class)
        .where(Class.id == class_id)
        .options(
            selectinload(Class.level),
            selectinload(Class.room),
        )
    )
    result = await db.execute(stmt)
    klass = result.scalar_one_or_none()
    if klass is None:
        raise NotFoundError("Class", class_id)
    return klass


async def _current_academic_year(db: AsyncSession) -> AcademicYear | None:
    stmt = select(AcademicYear).where(AcademicYear.is_current.is_(True)).limit(1)
    result = await db.execute(stmt)
    return result.scalar_one_or_none()


async def _load_students(db: AsyncSession, class_id: int, academic_year_id: int) -> list[dict]:
    """Charge les étudiants inscrits VALIDES dans la classe pour l'AY courante."""
    stmt = (
        select(Enrollment)
        .where(
            Enrollment.class_id == class_id,
            Enrollment.academic_year_id == academic_year_id,
            Enrollment.status == EnrollmentStatus.VALIDE.value,
        )
        .options(selectinload(Enrollment.student))
        .order_by(Enrollment.id.asc())
    )
    result = await db.execute(stmt)
    enrollments = list(result.scalars().all())

    # Pour chaque étudiant, récupérer le tel parent urgence (1er parent lié)
    student_ids = [e.student_id for e in enrollments if e.student]
    parent_phone_by_student: dict[int, str] = {}
    if student_ids:
        ps_stmt = (
            select(ParentStudent)
            .where(ParentStudent.student_id.in_(student_ids))
            .options(selectinload(ParentStudent.parent))
            .order_by(ParentStudent.student_id, ParentStudent.parent_id)
        )
        ps_rows = (await db.execute(ps_stmt)).scalars().all()
        for ps in ps_rows:
            sid = ps.student_id
            if sid not in parent_phone_by_student and ps.parent and ps.parent.phone:
                parent_phone_by_student[sid] = ps.parent.phone

    rows: list[dict] = []
    for e in enrollments:
        s = e.student
        if s is None:
            continue
        initials = f"{(s.first_name or '?')[:1]}{(s.last_name or '?')[:1]}".upper()
        rows.append(
            {
                "enrollment_number": getattr(s, "enrollment_number", None) or "",
                "first_name": s.first_name,
                "last_name": s.last_name,
                "genre": getattr(s, "genre", None),
                "birth_date": getattr(s, "birth_date", None),
                "initials": initials,
                "photo_url": getattr(s, "photo_url", None),
                "parent_phone": parent_phone_by_student.get(s.id),
                "birth_place": getattr(s, "birth_place", None),
                "nationality": getattr(s, "nationality", None),
                "is_repeater": e.is_repeater,
                "assignment_status": getattr(e.assignment_status, "value", e.assignment_status),
            }
        )

    # Tri par nom de famille puis prénom
    rows.sort(key=lambda r: ((r["last_name"] or ""), (r["first_name"] or "")))
    return rows


async def get_class_roster_pdf(db: AsyncSession, class_id: int) -> bytes:
    """Génère le PDF liste de classe pour une classe (AY courante)."""
    klass = await _load_class_with_room(db, class_id)
    ay = await _current_academic_year(db)
    if ay is None:
        raise NotFoundError("AcademicYear (current)", 0)

    students = await _load_students(db, class_id, ay.id)
    boys = sum(1 for s in students if s.get("genre") == "M")
    girls = sum(1 for s in students if s.get("genre") == "F")

    data = {
        "class_name": klass.name,
        "level_name": getattr(klass.level, "name", "") if klass.level else "",
        "academic_year_name": ay.name,
        "main_teacher_name": None,  # TODO : Class.main_teacher_id si ajouté plus tard
        "room_name": getattr(klass.room, "name", None) if klass.room else None,
        "students": students,
        "counts": {"total": len(students), "boys": boys, "girls": girls},
        "issued_at": datetime.utcnow(),
    }
    school = await _get_school_settings_dict(db)
    return generate_class_roster_pdf(data, school)


def _word_roster_document(data: dict) -> bytes:
    """Construire une véritable liste de classe éditable au format DOCX."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    from app.services.pdf._helpers import image_bytes
    from app.services.pdf.theme import PDFTheme

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.left_margin = section.right_margin = Cm(1.1)
    section.top_margin = section.bottom_margin = Cm(1.5)

    school = data.get("school_settings") or {}
    theme = PDFTheme.from_school(school)
    primary = RGBColor.from_string(theme.primary.lstrip("#").upper())
    accent = RGBColor.from_string(theme.accent.lstrip("#").upper())
    identity = doc.add_table(rows=1, cols=2)
    identity.autofit = False
    identity.columns[0].width = Cm(3)
    identity.columns[1].width = Cm(23)
    logo = image_bytes(school.get("logo_url"))
    if logo and logo[1] in ("image/png", "image/jpeg", "image/jpg"):
        try:
            identity.cell(0, 0).paragraphs[0].add_run().add_picture(BytesIO(logo[0]), width=Cm(2.4))
        except (ValueError, OSError):
            pass
    details_cell = identity.cell(0, 1)
    p = details_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name = p.add_run(school.get("school_name") or "Établissement")
    name.bold = True
    name.font.size = Pt(16)
    name.font.color.rgb = primary
    settings_lines = [
        school.get("drena_name"),
        (
            "Code établissement : " + str(school["ministry_code"])
            if school.get("ministry_code")
            else None
        ),
        school.get("address"),
        " · ".join(str(x) for x in (school.get("phone"), school.get("email")) if x),
        school.get("website"),
        school.get("motto"),
    ]
    for detail in filter(None, settings_lines):
        p = details_cell.add_paragraph(str(detail))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(8)

    title = doc.add_heading("LISTE DE CLASSE", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = primary
    doc.add_paragraph(
        f"Classe : {data['class_name']}  |  Année scolaire : {data['academic_year_name']}  |  Effectif : {len(data['students'])}"
    )
    headers = [
        "N°",
        "Matricule",
        "Nom",
        "Prénoms",
        "Sexe",
        "Date de naissance",
        "Lieu de naissance",
        "Nationalité",
        "Qualité",
        "Statut",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, headers, strict=True):
        cell.text = label
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = accent

    qualities = {True: "Redoublant", False: "Non redoublant"}
    statuses = {
        "affecte": "Affecté",
        "reaffecte": "Réaffecté",
        "non_affecte": "Non affecté",
    }
    for index, student in enumerate(data["students"], 1):
        born = student.get("birth_date")
        birth = born.strftime("%d/%m/%Y") if born else ""
        genre = getattr(student.get("genre"), "value", student.get("genre"))
        values = [
            str(index),
            student.get("enrollment_number") or "",
            student.get("last_name") or "",
            student.get("first_name") or "",
            genre or "",
            birth,
            student.get("birth_place") or "",
            student.get("nationality") or "",
            qualities.get(student.get("is_repeater"), ""),
            statuses.get(student.get("assignment_status"), ""),
        ]
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)
    doc.add_paragraph("Liste établie à partir des inscriptions validées de l'année scolaire courante.")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


async def get_class_roster_docx(db: AsyncSession, class_id: int) -> bytes:
    """La même population d'élèves que la liste de classe PDF, au format Word."""
    klass = await _load_class_with_room(db, class_id)
    ay = await _current_academic_year(db)
    if ay is None:
        raise NotFoundError("AcademicYear (current)", 0)
    students = await _load_students(db, class_id, ay.id)
    school = await _get_school_settings_dict(db)
    return _word_roster_document(
        {
            "class_name": klass.name,
            "academic_year_name": ay.name,
            "students": students,
            "school_settings": school,
        }
    )
