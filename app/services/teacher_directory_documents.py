"""Exports enseignants PDF/Word/Excel : mêmes données, colonnes et groupes.

L'identité documentaire provient UNIQUEMENT de school_settings du tenant ;
aucune couleur, logo, ville ou établissement n'est codé dans les exports.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from app.schemas.teacher_directory import TeacherExportOptions, page_orientation
from app.services.pdf import components as ui
from app.services.pdf._helpers import image_bytes
from app.services.pdf.theme import PDFTheme


def render_pdf(directory: dict[str, Any], school: dict[str, Any], opts: TeacherExportOptions) -> bytes:
    from weasyprint import HTML

    theme = PDFTheme.from_school(school)
    orientation = page_orientation(opts)
    page = "A4 landscape" if orientation == "landscape" else "A4"
    sections = []
    for group in directory["groups"]:
        # Une rubrique est une donnée documentaire, pas du HTML de confiance.
        title = ui.esc(group["title"])
        section = f'<div class="teacher-group"><h2>{title} <small>({len(group["rows"])})</small></h2>'
        section += ui.premium_table(
            headers=directory["headers"],
            rows=group["rows"],
            theme=theme,
            empty_message="Aucune affectation dans cette rubrique.",
        )
        sections.append(section + "</div>")

    if not sections:
        sections = ['<p class="no-results">Aucun enseignant ne correspond aux critères sélectionnés.</p>']

    year = ui.esc(directory["year"])
    summary = f'{directory["teachers_count"]} enseignant(s) · {directory["count"]} ligne(s) · {len(directory["groups"])} rubrique(s)'
    html = f"""<!DOCTYPE html><html lang="fr"><head><meta charset="UTF-8">
    {ui.base_styles(theme, page_size=page, margin="13mm 13mm", paginate=True)}
    <style>
       .teacher-group {{ margin: 11px 0 16px; break-inside: auto; }}
       .teacher-group h2 {{ font-size: 11px; color: var(--primary);
           border-bottom: 1px solid var(--border); padding: 3px 0 6px; margin-bottom: 6px;
           break-after: avoid; }}
       .teacher-group small {{ font-size: 9px; color: var(--muted); font-weight: normal; }}
       .teacher-group table {{ table-layout: fixed; overflow-wrap: anywhere; }}
       .teacher-group tr {{ break-inside: avoid; }}
       .no-results {{ color: var(--muted); text-align: center; padding: 30px 0; }}
    </style></head><body>
    {ui.premium_header(
        school, theme=theme, doc_type="LISTE DES ENSEIGNANTS",
        doc_subtitle=f"Année scolaire : {directory['year']} — {directory['group_by']}",
    )}
    {ui.meta_banner(
        f"<strong>Année scolaire :</strong> {year}",
        f"<strong>Effectif :</strong> {summary}",
        theme=theme,
    )}
    {''.join(sections)}
    {ui.premium_footer(school, theme=theme,
        note="Répertoire administratif — les affectations au référentiel ne présument pas d'une classe enseignée.")}
    </body></html>"""
    return HTML(string=html).write_pdf()


def render_word(directory: dict[str, Any], school: dict[str, Any], opts: TeacherExportOptions) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    orientation = page_orientation(opts)
    theme = PDFTheme.from_school(school)
    color = RGBColor.from_string(theme.primary.lstrip("#").upper())
    accent = RGBColor.from_string(theme.accent.lstrip("#").upper())
    doc = Document()
    section = doc.sections[0]
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.2)
    section.top_margin = section.bottom_margin = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(9)

    identity = doc.add_table(rows=1, cols=2)
    identity.columns[0].width = Cm(2.6)
    identity.columns[1].width = Cm(15 if orientation == "portrait" else 22)
    logo = image_bytes(school.get("logo_url"))
    if logo and logo[1] in ("image/png", "image/jpeg", "image/jpg"):
        try:
            identity.cell(0, 0).paragraphs[0].add_run().add_picture(
                BytesIO(logo[0]), width=Cm(2.0)
            )
        except (ValueError, OSError):
            pass
    details = identity.cell(0, 1)
    p = details.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name = p.add_run(school.get("school_name") or "Établissement")
    name.font.color.rgb = color
    name.font.size = Pt(15)
    name.bold = True
    for line in [
        school.get("drena_name"),
        f"Code MENA : {school['ministry_code']}" if school.get("ministry_code") else None,
        school.get("address"),
        " · ".join(str(school[k]) for k in ("phone", "email") if school.get(k)),
        school.get("website"),
        school.get("motto"),
    ]:
        if line:
            p = details.add_paragraph(str(line))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8)

    doc.add_heading("LISTE DES ENSEIGNANTS", level=0)
    doc.add_paragraph(
        f"Année : {directory['year']}  |  {directory['teachers_count']} enseignants  |  "
        f"{directory['count']} lignes  |  {directory['group_by']}"
    )
    for group in directory["groups"]:
        title = doc.add_heading(f"{group['title']} ({len(group['rows'])})", level=2)
        for run in title.runs:
            run.font.color.rgb = color
        table = doc.add_table(rows=1, cols=len(directory["headers"]))
        table.style = "Table Grid"
        for cell, heading in zip(table.rows[0].cells, directory["headers"], strict=True):
            cell.text = heading
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(8)
                run.font.color.rgb = accent
        for values in group["rows"]:
            cells = table.add_row().cells
            for cell, value in zip(cells, values, strict=True):
                cell.text = str(value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(8)
    if not directory["groups"]:
        doc.add_paragraph("Aucun enseignant ne correspond aux critères.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(school.get("school_name") or "").font.color.rgb = color
    footer.add_run(" · Document administratif interne")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _safe_excel(value: object) -> str:
    """Empêche l'exécution des formules injectées depuis un nom ou une matière."""
    text = str(value if value is not None else "")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def render_excel(directory: dict[str, Any], school: dict[str, Any], opts: TeacherExportOptions) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    from app.services.exports._workbook_branding import (
        primary_argb, set_widths, setup_printing, style_table_header, write_header,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Enseignants"
    headers = directory["headers"]
    row_index = write_header(
        ws, school, title="LISTE DES ENSEIGNANTS",
        subtitle=f"Année {directory['year']} — {directory['group_by']}",
        meta_lines=[
            f"{directory['teachers_count']} enseignants · {directory['count']} lignes",
            "Affectations référentiel : sans classe et sans rattachement à l'année scolaire.",
        ],
        width=len(headers),
    )
    header_row = row_index
    for idx, label in enumerate(headers, 1):
        ws.cell(row_index, idx, label)
    style_table_header(ws, row_index, len(headers), school)
    row_index += 1
    for group in directory["groups"]:
        ws.merge_cells(start_row=row_index, start_column=1,
                       end_row=row_index, end_column=max(1, len(headers)))
        group_cell = ws.cell(row_index, 1, _safe_excel(f"{group['title']} — {len(group['rows'])} ligne(s)"))
        group_cell.font = Font(bold=True, color=primary_argb(school))
        row_index += 1
        for values in group["rows"]:
            for idx, value in enumerate(values, 1):
                ws.cell(row_index, idx, _safe_excel(value))
                ws.cell(row_index, idx).alignment = Alignment(vertical="center", wrap_text=True)
            row_index += 1
    if not directory["groups"]:
        ws.cell(row_index, 1, "Aucun enseignant ne correspond aux critères.")
    set_widths(ws, [min(62, max(18, len(label) + 10)) for label in headers])
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, row_index - 1)}" if not directory["groups"] else None
    setup_printing(ws, header_row=header_row, landscape=page_orientation(opts) == "landscape")
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
